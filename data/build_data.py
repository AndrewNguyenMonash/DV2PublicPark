"""Build the DV2 combined/derived datasets from the raw council, ABS and AusPlay extracts.

Run from the DV2/data folder:  python3 build_data.py
Inputs : raw/  (extracts downloaded from each publisher on 13 Oct 2026)
Outputs: derived/
"""
import csv, io, json, math, zipfile, datetime as dt
from collections import defaultdict, Counter

import h3
import openpyxl
import shapefile
import topojson as tp
from pyproj import Transformer
from shapely.geometry import shape, mapping, Point, Polygon, MultiPolygon
from shapely.ops import unary_union, transform

RAW, OUT = "raw/", "derived/"
TO_ALBERS = Transformer.from_crs("EPSG:4326", "EPSG:3577", always_xy=True).transform   # GDA94 Australian Albers (equal area)
ALBERS_TO_LL = Transformer.from_crs("EPSG:3577", "EPSG:4326", always_xy=True).transform   # back to lon/lat
MGA54_TO_LL = Transformer.from_crs("EPSG:28354", "EPSG:4326", always_xy=True).transform
WEBM_TO_LL = Transformer.from_crs("EPSG:3857", "EPSG:4326", always_xy=True).transform

COUNCILS = {  # code: (display name, ABS LGA 2025 name, state)
    "BNE": ("Brisbane City Council", "Brisbane", "QLD"),
    "SYD": ("City of Sydney", "Sydney", "NSW"),
    "MEL": ("City of Melbourne", "Melbourne", "VIC"),
    "PER": ("City of Perth", "Perth", "WA"),
    "ADL": ("City of Adelaide", "Adelaide", "SA"),
    "HBA": ("City of Hobart", "Hobart", "TAS"),
    "CBR": ("ACT Government (City Services)", "Unincorporated ACT", "ACT"),
    "DRW": ("City of Darwin", "Darwin", "NT"),
}
TYPES = {"P": "Playground", "B": "BBQ", "F": "Drinking fountain", "T": "Toilet in a park", "D": "Dog off-leash area"}
SIZE_CLASSES = [(0.5, "Pocket (under 0.5 ha)"), (2, "Local (0.5-2 ha)"), (10, "District (2-10 ha)"), (1e9, "Regional (10 ha+)")]


def size_class(ha):
    for lim, name in SIZE_CLASSES:
        if ha < lim:
            return name


def albers_ha(geom):
    return transform(TO_ALBERS, geom).area / 1e4


def read_csv(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def write_csv(path, rows, fields):
    with open(OUT + path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        for r in rows:
            w.writerow({k: r.get(k, "") for k in fields})
    print(f"  wrote {path}: {len(rows)} rows")


def dedupe(polys):
    """Drop polygons that are mostly (>50 %) covered by a larger polygon already kept (duplicate versions)."""
    kept, union = [], None
    for item in sorted(polys, key=lambda t: -t[1].area):
        g = item[1]
        if union is not None and g.area > 0 and g.intersection(union).area / g.area > 0.5:
            continue
        kept.append(item)
        union = g if union is None else union.union(g)
    return kept


# ---------------------------------------------------------------- 1. council parks
print("1. council parks")
parks = []  # dicts: council, park_id, park_name, area_ha, lon, lat
park_geoms = {}  # council -> unioned shapely geometry (lon/lat) for 'in park' tests

for r in read_csv(RAW + "bne_parks.csv"):
    parks.append(dict(council="BNE", park_id=r["park_id"], park_name=r["park_name"], area_ha=float(r["area_ha"]),
                      lon=float(r["lon"]), lat=float(r["lat"])))


def from_geojson(code, path, name_key="park_name", do_dedupe=True):
    fc = json.load(open(path))
    items = []
    for f in fc["features"]:
        if not f.get("geometry"):
            continue
        g = shape(f["geometry"]).buffer(0)
        if g.is_empty:
            continue
        items.append((f["properties"], g))
    if do_dedupe:
        items = dedupe(items)
    for props, g in items:
        c = g.representative_point()
        parks.append(dict(council=code, park_id=props.get("park_id"), park_name=(props.get(name_key) or "").strip(),
                          park_type=props.get("park_type") or props.get("category") or "",
                          area_ha=round(albers_ha(g), 3), lon=round(c.x, 5), lat=round(c.y, 5)))
    park_geoms[code] = unary_union([g for _, g in items])
    print(f"  {code}: {len(fc['features'])} features -> {len(items)} after removing duplicates")


from_geojson("SYD", RAW + "syd_parks.geojson")
from_geojson("MEL", RAW + "mel_parks.geojson")
from_geojson("PER", RAW + "per_parks.geojson")
from_geojson("HBA", RAW + "hba_parks.geojson")
from_geojson("CBR", RAW + "cbr_parks_ext.geojson", do_dedupe=False)   # incl. pedestrian parkland + semi-natural open space

# Adelaide Park Lands parks (shapefile, WGS84)
zf = zipfile.ZipFile(RAW + "adl_parks_shp.zip")
rd = shapefile.Reader(shp=io.BytesIO(zf.read("ZIP/Parks.shp")), dbf=io.BytesIO(zf.read("ZIP/Parks.dbf")),
                      shx=io.BytesIO(zf.read("ZIP/Parks.shx")))
items = []
for sr in rd.iterShapeRecords():
    g = shape(sr.shape.__geo_interface__).buffer(0)
    rec = sr.record.as_dict()
    items.append(({"park_id": rec.get("PARKNUM"), "park_name": rec.get("PARKNAME")}, g))
items = dedupe(items)
for props, g in items:
    c = g.representative_point()
    parks.append(dict(council="ADL", park_id=props["park_id"], park_name=props["park_name"], area_ha=round(albers_ha(g), 3),
                      lon=round(c.x, 5), lat=round(c.y, 5)))
park_geoms["ADL"] = unary_union([g for _, g in items])
print(f"  ADL: {len(rd)} features -> {len(items)}")

for p in parks:
    p["size_class"] = size_class(p["area_ha"])
write_csv("council_parks.csv", parks, ["council", "park_id", "park_name", "park_type", "area_ha", "size_class", "lon", "lat"])

# ---------------------------------------------------------------- 2. facilities (all councils)
print("2. facilities")
fac = []


def add(code, t, lon, lat, unit="asset", note=""):
    fac.append(dict(council=code, type=t, type_label=TYPES[t], lon=round(float(lon), 5), lat=round(float(lat), 5), unit=unit, note=note))


for code in ["BNE", "SYD", "MEL", "PER", "HBA", "CBR", "DRW"]:
    for r in read_csv(RAW + f"{code.lower()}_facilities.csv"):
        add(code, r["type"], r["lon"], r["lat"], r.get("unit") or "asset")

# Adelaide: BBQs + drinking fountains (MGA zone 54), toilets (lon/lat), playgrounds (polygons -> centroid)
def adl_ll(r):
    e, n = r["Easting (Location/Map Coordinates)"], r["Northing (Location/Map Coordinates)"]
    if e and n:
        return MGA54_TO_LL(float(e), float(n))
    g = (r.get("Geometry") or "").strip()
    if g.startswith("POINT"):
        x, y = g[g.index("(") + 1:g.index(")")].split()
        return WEBM_TO_LL(float(x), float(y))
    return None


adl_skipped = 0
for path, t in [("adl_bbqs_raw.csv", "B"), ("adl_drinking_fountains_raw.csv", "F")]:
    for r in read_csv(RAW + path):
        ll = adl_ll(r)
        if ll is None:
            adl_skipped += 1
            continue
        add("ADL", t, *ll)
print("  Adelaide rows without coordinates skipped:", adl_skipped)
for r in read_csv(RAW + "adl_toilets_raw.csv"):
    add("ADL", "T", r["POINT_X"], r["POINT_Y"])
zf = zipfile.ZipFile(RAW + "adl_playgrounds_shp.zip")
rd = shapefile.Reader(shp=io.BytesIO(zf.read("ZIP/Playgrounds.shp")), dbf=io.BytesIO(zf.read("ZIP/Playgrounds.dbf")),
                      shx=io.BytesIO(zf.read("ZIP/Playgrounds.shx")))
for s in rd.shapes():
    c = shape(s.__geo_interface__).representative_point()
    add("ADL", "P", c.x, c.y)

# Keep only drinking fountains and toilets that sit in (or within ~30 m of) a council park, where park outlines exist.
# Brisbane's BBQs/fountains/dog areas are already park assets and its toilets were filtered to park polygons on download.
# ACT (Place Management urban open space assets) and Darwin (no park outlines) are kept as published.
buffered = {k: transform(TO_ALBERS, v).buffer(30) for k, v in park_geoms.items()}
kept = []
dropped = Counter()
for f in fac:
    if f["type"] in ("F", "T") and f["council"] in buffered and f["council"] != "CBR":
        x, y = TO_ALBERS(f["lon"], f["lat"])
        if not buffered[f["council"]].contains(Point(x, y)):
            dropped[(f["council"], f["type"])] += 1
            continue
        f["note"] = "within 30 m of a council park"
    kept.append(f)
fac = kept
print("  dropped (not in a council park):", dict(dropped))

# Council outlines (ABS LGA 2025) in metres (Albers): used to drop facilities outside their council (below)
# and, back in lon/lat, as the C6 base maps (section 3).
# ABS draws council areas on land only, so the tidal Brisbane River (about 200 m wide) is left out and splits
# Brisbane into a north half and a south half. Growing the outline by 200 m and shrinking it back by 200 m
# closes every gap narrower than 400 m; the wider stretches of river, now enclosed by the city, are holes that
# are then filled. Moreton Island, the smaller bay islands and the river mouth stay as they are.
# The other seven outlines are already single shapes.
lga = json.load(open(RAW + "abs_lga2025_8councils.geojson"))
by_name = {f["properties"]["lga_name_2025"]: f for f in lga["features"]}
outlines = {}
for code, (cname, lga_name, st) in COUNCILS.items():
    g = transform(TO_ALBERS, shape(by_name[lga_name]["geometry"]))
    if code == "BNE":
        g = g.buffer(200).buffer(-200)
        g = MultiPolygon([Polygon(p.exterior) for p in g.geoms])   # keep each piece's outer edge only (fills the holes)
    outlines[code] = g

# Drop facilities more than 400 m outside their council's outline. 400 m allows for the simplified ABS coastline
# (the fountains on Brisbane's Shorncliffe Pier are 330 m out). It drops one City of Sydney dog off-leash area
# in Centennial Park, a state-run parkland 510 m outside the city.
near = {code: g.buffer(400) for code, g in outlines.items()}
kept, outside = [], Counter()
for f in fac:
    x, y = TO_ALBERS(f["lon"], f["lat"])
    if near[f["council"]].contains(Point(x, y)):
        kept.append(f)
    else:
        outside[(f["council"], f["type"])] += 1
fac = kept
print("  dropped (outside the council):", dict(outside))
write_csv("facilities_all_councils.csv", fac, ["council", "type", "type_label", "lon", "lat", "unit", "note"])

# ---------------------------------------------------------------- 3. council summary + boundaries
print("3. council summary")
counts = Counter((f["council"], f["type"]) for f in fac)
published = {  # facility types each council publishes (False = no open dataset found)
    "BNE": dict(P=False, B=True, F=True, T=True, D=True),
    "SYD": dict(P=True, B=False, F=True, T=True, D=True),
    "MEL": dict(P=True, B=True, F=True, T=True, D=False),
    "PER": dict(P=True, B=True, F=True, T=True, D=False),
    "ADL": dict(P=True, B=True, F=True, T=True, D=False),
    "HBA": dict(P=True, B=True, F=True, T=True, D=True),
    "CBR": dict(P=True, B=True, F=True, T=True, D=True),
    "DRW": dict(P=True, B=True, F=True, T=True, D=False),
}
facility_years = {  # year(s) of each council's facility data, from the dataset dates in README.md (S2-S12)
    "BNE": dict(B="2026", F="2026", T="2026", D="2026"),
    "SYD": dict(P="2024-25", F="2024-25", T="2024-25", D="2024-25"),
    "MEL": dict(P="2021", B="2026", F="2021", T="2021"),
    "PER": dict(P="2022", B="2026", F="2026", T="2022"),
    "ADL": dict(P="2013-16", B="2013-16", F="2013-16", T="2013-16"),
    "HBA": dict(P="2025", B="2025", F="2025", T="2025", D="2025"),
    "CBR": dict(P="2026", B="2026", F="2026", T="2026", D="2026"),
    "DRW": dict(P="2024-26", B="2024-26", F="2024-26", T="2024-26"),
}
summary, bnd_features = [], []
for code, (cname, lga_name, st) in COUNCILS.items():
    f = by_name[lga_name]
    pr = f["properties"]
    g = shape(f["geometry"])
    rp = g.representative_point()
    cp = [p for p in parks if p["council"] == code]
    erp = pr["erp_2025"]
    park_ha = round(albers_ha(park_geoms[code]), 1) if code in park_geoms else (round(sum(p["area_ha"] for p in cp), 1) if cp else None)
    row = dict(council=code, council_name=cname, state=st, lga_code_2025=pr["lga_code_2025"], lga_name_2025=lga_name,
               erp_2025=erp, lga_area_km2=round(pr["area_km2"], 1), lon=round(rp.x, 4), lat=round(rp.y, 4),
               n_parks=len(cp) if cp else "", park_area_ha=park_ha if park_ha is not None else "",
               park_m2_per_resident=round(park_ha * 1e4 / erp, 1) if park_ha else "",
               park_share_of_council_area_pct=round(park_ha / (pr["area_km2"] * 100) * 100, 1) if park_ha else "")
    for t, label in [("P", "playgrounds"), ("B", "bbqs"), ("F", "drinking_fountains"), ("T", "park_toilets"), ("D", "dog_offleash_areas")]:
        n = counts.get((code, t), 0)
        row[label] = n if published[code][t] else ""
        row[label + "_per_10k"] = round(n / erp * 1e4, 2) if published[code][t] else ""
    summary.append(row)
    # Simplify by 20 m (in metres, before going back to lon/lat): it removes the extra corner points the Brisbane
    # buffers add and is under half a pixel even on the smallest C6 map (Adelaide, about 45 m per pixel).
    bnd_features.append({"type": "Feature", "properties": {"council": code, "council_name": cname, "state": st},
                         "geometry": mapping(transform(ALBERS_TO_LL, outlines[code].simplify(20)))})
fields = list(summary[0].keys())
write_csv("council_summary.csv", summary, fields)

# Long format for the heatmap (one row per council x facility type)
long_rows = []
for s in summary:
    for t, label in [("P", "playgrounds"), ("B", "bbqs"), ("F", "drinking_fountains"), ("T", "park_toilets"), ("D", "dog_offleash_areas")]:
        long_rows.append(dict(council=s["council"], council_name=s["council_name"], facility=TYPES[t],
                              count=s[label], per_10k_residents=s[label + "_per_10k"],
                              per_100_parks=(round(s[label] / s["n_parks"] * 100, 1) if s[label] != "" and s["n_parks"] else ""),
                              published="yes" if s[label] != "" else "no",
                              data_year=facility_years[s["council"]].get(t, "")))
write_csv("council_facility_rates_long.csv", long_rows, ["council", "council_name", "facility", "count", "per_10k_residents", "per_100_parks", "published", "data_year"])

# Park size classes (treemap: council > size class)
sz = defaultdict(lambda: [0, 0.0])
for p in parks:
    sz[(p["council"], p["size_class"])][0] += 1
    sz[(p["council"], p["size_class"])][1] += p["area_ha"]
write_csv("park_size_classes.csv",
          [dict(council=k[0], size_class=k[1], n_parks=v[0], area_ha=round(v[1], 1)) for k, v in sorted(sz.items())],
          ["council", "size_class", "n_parks", "area_ha"])

# Boundaries as TopoJSON (object 'councils')
# prequantize=1e6 rounds points to about 4 m. 1e5 (about 37 m) made the Sydney outline cross itself at the Circular Quay wharves.
topo = tp.Topology({"type": "FeatureCollection", "features": bnd_features}, object_name="councils", prequantize=1e6)
open(OUT + "council_boundaries.topojson", "w").write(topo.to_json())
print("  wrote council_boundaries.topojson")

# Australian states basemap: drop islands under 300 km2, simplify.
# The ABS file mixes up outlines and holes for NSW: the NSW mainland is stored as a hole and the
# ACT (a hole in NSW) as the outline, so reading the file as it is turned NSW into a copy of the ACT.
# Each state is therefore rebuilt from its rings by where they sit: rings under 300 km2 are dropped,
# a ring inside another ring of the same state is a hole (the ACT inside NSW), every other ring is land.
ste = json.load(open(RAW + "abs_ste2026_simplified.geojson"))
sf = []
for f in ste["features"]:
    geom = f["geometry"]
    polys = geom["coordinates"] if geom["type"] == "MultiPolygon" else [geom["coordinates"]]
    rings = [Polygon(r).buffer(0) for poly in polys for r in poly if len(r) >= 4]
    rings = [r for r in rings if transform(TO_ALBERS, r).area > 3e8]
    land, holes = [], []
    for r in rings:
        inside = any(o.contains(r.representative_point()) for o in rings if o is not r)
        (holes if inside else land).append(r)
    g2 = unary_union(land).difference(unary_union(holes)).simplify(0.04, preserve_topology=True)
    sf.append({"type": "Feature", "properties": {"state": f["properties"]["STATE_NAME_2026"]}, "geometry": mapping(g2)})
topo = tp.Topology({"type": "FeatureCollection", "features": sf}, object_name="states", prequantize=1e4, toposimplify=0.02)
open(OUT + "australia_states.topojson", "w").write(topo.to_json())
print("  wrote australia_states.topojson")

# ---------------------------------------------------------------- 4. Brisbane suburbs (choropleth)
print("4. Brisbane suburbs")
bne_geom = shape(by_name["Brisbane"]["geometry"])
bne_land = bne_geom.buffer(0)   # the same outline made valid, so suburbs can be cut along it
sal = json.load(open(RAW + "abs_sal2021_brisbane_area.geojson"))
sub_parks = {r["suburb"].upper(): r for r in read_csv(RAW + "bne_suburb_parks.csv")}
rows, feats = [], []
for f in sal["features"]:
    g = shape(f["geometry"]).buffer(0)
    # Keep a suburb if at least 10% of its land is in Brisbane (slivers along the boundary are dropped).
    # Four localities are split by the council boundary (Chuwar, Lake Manchester, England Creek, Banks Creek):
    # they are cut at the boundary and their area is measured again, so their park share is a share of the
    # Brisbane part only. Every other suburb is at least 95% inside and keeps its ABS outline and area.
    inside = g.intersection(bne_land)
    share_inside = albers_ha(inside) / albers_ha(g)
    if share_inside < 0.10:
        continue
    nm = f["properties"]["sal_name_2021"]
    key = nm.split(" (")[0].upper()          # 'Ascot (Brisbane - Qld)' -> 'ASCOT'
    sp = sub_parks.get(key)
    area_ha = f["properties"]["area_albers_sqkm"] * 100
    if share_inside < 0.95:
        g = inside
        area_ha = albers_ha(g)
    park_ha = float(sp["park_ha"]) if sp else 0.0
    row = dict(sal_code_2021=f["properties"]["sal_code_2021"], suburb=nm, suburb_area_ha=round(area_ha, 1),
               n_parks=int(sp["n_parks"]) if sp else 0, council_park_ha=round(park_ha, 1),
               pct_council_park=round(min(park_ha / area_ha * 100, 100), 1) if area_ha else "")
    rows.append(row)
    feats.append({"type": "Feature", "properties": {"sal_code_2021": row["sal_code_2021"], "suburb": nm},
                  "geometry": mapping(g.simplify(0.0006, preserve_topology=True))})
matched = sum(1 for r in rows if r["n_parks"])
write_csv("bne_suburb_park_share.csv", rows, list(rows[0].keys()))
print(f"  suburbs in Brisbane LGA: {len(rows)}; with matched council parks: {matched}; park-suburb names unmatched:",
      sorted(set(sub_parks) - {r['suburb'].split(' (')[0].upper() for r in rows})[:10])
topo = tp.Topology({"type": "FeatureCollection", "features": feats}, object_name="suburbs", prequantize=1e5)
open(OUT + "bne_suburbs.topojson", "w").write(topo.to_json())
print("  wrote bne_suburbs.topojson")

# ---------------------------------------------------------------- 5. Brisbane hexbins (H3 resolution 8)
print("5. Brisbane hexbins")
# Mainland Brisbane only (the council outline from section 3, with the river filled in), so C8 covers the same
# area as C7. Moreton Island and the bay islands would add 221 empty hexagons for one facility (a toilet).
mainland_m = max(outlines["BNE"].geoms, key=lambda p: p.area)        # in metres
hx = defaultdict(Counter)
for f in fac:
    if f["council"] == "BNE" and mainland_m.buffer(500).contains(Point(*TO_ALBERS(f["lon"], f["lat"]))):
        hx[h3.latlng_to_cell(f["lat"], f["lon"], 8)][f["type"]] += 1
# cover the whole mainland so empty hexagons are drawn too
cover = h3.geo_to_cells(mapping(transform(ALBERS_TO_LL, mainland_m)), 8)
# Each hexagon is named after the suburb its centre falls in (the nearest suburb if the centre is on the river).
sub_shapes = [(ft["properties"]["suburb"].split(" (")[0], shape(ft["geometry"])) for ft in feats]
cells = set(cover) | set(hx)
# Add any hexagon enclosed on all six sides (one, near Mount Crosby, where a bend of the river pokes into the outline).
cells |= {n for c in cells for n in h3.grid_disk(c, 1)
          if n not in cells and all(m in cells for m in h3.grid_disk(n, 1) if m != n)}
hex_feats = []
for cell in sorted(cells):   # sorted, so the file comes out the same on every run
    ring = [(lng, lat) for lat, lng in h3.cell_to_boundary(cell)]
    ring.append(ring[0])
    c = hx.get(cell, Counter())
    lat, lng = h3.cell_to_latlng(cell)
    centre = Point(lng, lat)
    inside = [name for name, s in sub_shapes if s.contains(centre)]
    suburb = inside[0] if inside else min(sub_shapes, key=lambda t: t[1].distance(centre))[0]
    hex_feats.append({"type": "Feature", "properties": {"h3": cell, "suburb": suburb, "bbq": c["B"], "drinking_fountain": c["F"],
                                                        "dog_offleash": c["D"], "park_toilet": c["T"],
                                                        "total": sum(c.values())},
                      "geometry": {"type": "Polygon", "coordinates": [[[round(x, 5), round(y, 5)] for x, y in ring]]}})
topo = tp.Topology({"type": "FeatureCollection", "features": hex_feats}, object_name="hexbins", prequantize=1e5)
open(OUT + "bne_facility_hexbins.topojson", "w").write(topo.to_json())
print(f"  hexagons: {len(hex_feats)} (edge ~{h3.average_hexagon_edge_length(8, 'km'):.2f} km)")

# ---------------------------------------------------------------- 6. Brisbane bookings
print("6. Brisbane bookings")
cal = []
for r in read_csv(RAW + "bne_bookings_daily.csv"):
    d = dt.date.fromisoformat(r["date"])
    tot = int(r["commercial_activity_bookings"]) + int(r["park_booking_applications"])
    cal.append(dict(date=r["date"], year=d.year, month=d.month, iso_week=d.isocalendar()[1], weekday=d.strftime("%a"),
                    commercial_activity_bookings=r["commercial_activity_bookings"],
                    park_booking_applications=r["park_booking_applications"], total_bookings=tot,
                    note="booking system roll-out (counts not comparable)" if d < dt.date(2024, 1, 1) else ""))
write_csv("bne_bookings_calendar.csv", cal, list(cal[0].keys()))

hr = read_csv(RAW + "bne_bookings_hour_2024_2025.csv")
hr = [dict(day_type=r["day_type"], hour=int(r["hour"]), bookings_2024_2025=int(r["bookings_2024_2025"]),
           avg_bookings_per_day=float(r["avg_bookings_per_day"])) for r in hr]
have = {(r["day_type"], r["hour"]) for r in hr}
for d in ("weekday", "weekend"):            # fill empty hours with 0 so the radial chart has all 24 spokes
    for h in range(24):
        if (d, h) not in have:
            hr.append(dict(day_type=d, hour=h, bookings_2024_2025=0, avg_bookings_per_day=0.0))
hr.sort(key=lambda r: (r["day_type"], r["hour"]))
write_csv("bne_bookings_by_hour.csv", hr, ["day_type", "hour", "bookings_2024_2025", "avg_bookings_per_day"])

rk = read_csv(RAW + "bne_bookings_park_quarter_rank.csv")
for r in rk:                                # "MT COOT-THA RESERVE" -> "Mt Coot-tha Reserve", for the chart's line labels
    r["park_name"] = " ".join(w.capitalize() for w in r["park_name"].split())
write_csv("bne_bookings_top_parks_quarterly.csv", rk, ["quarter", "park_name", "bookings", "rank_among_all_parks"])

# ---------------------------------------------------------------- 7. AusPlay (calendar year 2025 tables)
print("7. AusPlay")
wb = openpyxl.load_workbook(RAW + "ausplay_national_2026-04-30.xlsx", read_only=True, data_only=True)


def ausplay_table(sheet, header_row, first_data_row, last_data_row):
    rows = list(wb[sheet].iter_rows(values_only=True))
    hdr = [str(c).replace("\n", " ").replace("  ", " ").strip() for c in rows[header_row] if c is not None]
    out = []
    for r in rows[first_data_row:last_data_row]:
        vals = [c for c in r if c is not None]
        nums = [v for v in vals if isinstance(v, (int, float))]
        labels = [str(v).strip() for v in vals if isinstance(v, str)]
        if not nums:
            continue
        age = labels[-1] if labels else ""
        out.append((age, dict(zip(hdr[-len(nums):], nums))))
    return out


places = ausplay_table("7", 12, 14, 22)   # 'All' block: Total + 7 age bands
pl_rows = []
for age, d in places:
    for k, v in d.items():
        if k == "Total":
            continue
        pl_rows.append(dict(age_group=age, place=k, participation_rate_pct=round(v * 100, 1)))
write_csv("ausplay_2025_places.csv", pl_rows, ["age_group", "place", "participation_rate_pct"])

mot = ausplay_table("12", 12, 14, 22)
mo_rows = []
for age, d in mot:
    for k, v in d.items():
        mo_rows.append(dict(age_group=age, motivation=k, share_of_active_adults_pct=round(v * 100, 1)))
write_csv("ausplay_2025_motivations.csv", mo_rows, ["age_group", "motivation", "share_of_active_adults_pct"])
print("done")
